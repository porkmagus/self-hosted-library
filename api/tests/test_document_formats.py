import fcntl
import subprocess
from pathlib import Path

import fitz
import pytest
from docx import Document
from ebooklib import epub

from api.services.parser_svc import (
    DocumentParseError,
    DocumentProcessingError,
    extract_docx,
    extract_epub,
    extract_legacy_doc,
    extract_pdf_ocr,
    extract_pdf_pymupdf,
    extract_text_file,
)

GOLDEN_TEXT = "The moon is a mirror; memory is the silver behind it."


def test_pdf_text_extraction_preserves_searchable_content(tmp_path: Path) -> None:
    path = tmp_path / "golden.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), GOLDEN_TEXT)
    document.save(path)
    document.close()

    assert GOLDEN_TEXT in extract_pdf_pymupdf(path)


def test_pdf_ocr_runs_isolated_with_one_cpu_thread(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "scan.pdf"
    source.write_bytes(b"scan")
    calls = []

    def run(command, **kwargs):
        calls.append((command, kwargs))
        Path(command[-1]).write_text(GOLDEN_TEXT)
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(subprocess, "run", run)

    assert extract_pdf_ocr(source, tmp_path) == GOLDEN_TEXT
    command, kwargs = calls[0]
    assert command[1:3] == ["-m", "api.services.pdf_ocr_worker"]
    assert kwargs["env"]["OMP_THREAD_LIMIT"] == "1"
    assert 0 < kwargs["timeout"] <= 3600
    assert int(kwargs["env"]["PDF_OCR_MAX_INPUT_BYTES"]) > 0
    assert int(kwargs["env"]["PDF_OCR_MAX_PAGES"]) > 0
    assert int(kwargs["env"]["PDF_OCR_MAX_PAGE_PIXELS"]) > 0
    assert int(kwargs["env"]["PDF_OCR_MAX_OUTPUT_BYTES"]) > 0
    assert int(kwargs["env"]["PDF_OCR_MAX_MEMORY_BYTES"]) > 0
    assert int(kwargs["env"]["PDF_OCR_MAX_CPU_SECONDS"]) > 0


def test_pdf_ocr_lock_wait_is_bounded(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "scan.pdf"
    source.write_bytes(b"scan")
    monkeypatch.setenv("PDF_OCR_LOCK_PATH", str(tmp_path / "ocr.lock"))
    monkeypatch.setenv("PDF_OCR_LOCK_TIMEOUT_SECONDS", "0.01")
    monkeypatch.setattr(
        fcntl,
        "flock",
        lambda *_args: (_ for _ in ()).throw(BlockingIOError()),
    )

    def unexpected_run(*_args, **_kwargs):
        raise AssertionError("OCR subprocess must not start without the lock")

    monkeypatch.setattr(subprocess, "run", unexpected_run)

    with pytest.raises(DocumentProcessingError, match="OCR lock acquisition timed out"):
        extract_pdf_ocr(source, tmp_path)


def test_pdf_ocr_rejects_non_finite_lock_timeout(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "scan.pdf"
    source.write_bytes(b"scan")
    monkeypatch.setenv("PDF_OCR_LOCK_TIMEOUT_SECONDS", "inf")

    with pytest.raises(DocumentProcessingError, match="must be finite"):
        extract_pdf_ocr(source, tmp_path)


@pytest.mark.parametrize(
    ("returncode", "expected_error"),
    [(42, DocumentParseError), (1, DocumentProcessingError)],
)
def test_native_pdf_worker_exit_classification(
    tmp_path: Path, monkeypatch, returncode: int, expected_error: type[Exception]
) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(b"bad-pdf")
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *_args, **_kwargs: subprocess.CompletedProcess(
            [], returncode, "", "classified failure"
        ),
    )

    with pytest.raises(expected_error, match="classified failure"):
        extract_pdf_pymupdf(source)


def test_docx_text_extraction_preserves_paragraph_order(tmp_path: Path) -> None:
    path = tmp_path / "golden.docx"
    document = Document()
    document.add_paragraph("First invocation")
    document.add_paragraph("Second invocation")
    document.save(str(path))

    assert extract_docx(path) == "First invocation\n\nSecond invocation"


def test_legacy_doc_extraction_uses_antiword(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"legacy-word")
    commands: list[list[str]] = []

    def run(command, **kwargs):
        commands.append(command)
        assert kwargs.get("capture_output") is not True
        kwargs["stdout"].write(b"First\nSecond\n")
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(subprocess, "run", run)

    assert extract_legacy_doc(path) == "First\nSecond"
    assert commands == [["antiword", "-m", "UTF-8.txt", str(path)]]


def test_mislabeled_rtf_doc_uses_unrtf(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes(br"{\rtf1 Golden text}")
    commands: list[list[str]] = []

    def run(command, **kwargs):
        commands.append(command)
        assert kwargs.get("capture_output") is not True
        kwargs["stdout"].write(b"Golden text\n")
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(subprocess, "run", run)

    assert extract_legacy_doc(path) == "Golden text"
    assert commands == [["unrtf", "--text", "--nopict", str(path)]]


def test_legacy_converter_timeout_is_systemic(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"legacy-word")
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda command, **kwargs: (_ for _ in ()).throw(
            subprocess.TimeoutExpired(command, kwargs["timeout"])
        ),
    )

    with pytest.raises(DocumentProcessingError, match="timed out"):
        extract_legacy_doc(path)


def test_mislabeled_plaintext_doc_uses_encoding_fallback(
    tmp_path: Path, monkeypatch
) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes("Café grimoire".encode("cp1252"))
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *_args, **_kwargs: pytest.fail("plaintext must not use a converter"),
    )

    assert extract_legacy_doc(path) == "Café grimoire"


def test_epub_text_extraction_reads_document_items(tmp_path: Path) -> None:
    path = tmp_path / "golden.epub"
    book = epub.EpubBook()
    book.set_identifier("golden")
    book.set_title("Golden")
    chapter = epub.EpubHtml(title="Chapter", file_name="chapter.xhtml")
    chapter.content = f"<html><body><p>{GOLDEN_TEXT}</p></body></html>"
    book.add_item(chapter)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.toc = (chapter,)
    book.spine = ["nav", chapter]
    epub.write_epub(path, book)

    assert GOLDEN_TEXT in extract_epub(path)


def test_text_extraction_falls_back_to_legacy_encoding(tmp_path: Path) -> None:
    path = tmp_path / "legacy.txt"
    path.write_bytes("Caf\u00e9 grimoires".encode("cp1252"))

    assert extract_text_file(path) == "Caf\u00e9 grimoires"
